from io import BytesIO
import json
from pathlib import Path
from uuid import uuid4
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject

from g3.app import create_app
from g3.config import Settings
from g3.input_documents import DocumentError, extract_document

SOURCE = "任务交接要说明交付目标、必要背景和完成标准。目标说明最终交付物，背景说明读者与已有材料，完成标准列出必须完成的内容。未记录的信息应标为待确认，不要猜测补齐。"
READY = lambda _: {"ready": True, "checks": {"diagram_layout": True}, "blockers": []}


def docx_bytes():
    d = Document();d.add_paragraph("任务交接材料")
    t = d.add_table(rows=1, cols=2);t.cell(0,0).text="目标";t.cell(0,1).text="说明交付物"
    d.add_paragraph(SOURCE)
    b=BytesIO();d.save(b);return b.getvalue()


def pdf_bytes(text="Document input test", encrypted=False, image=False):
    writer=PdfWriter();page=writer.add_blank_page(width=595,height=842)
    if image:
        obj=DecodedStreamObject();obj.set_data(b'\xff\xff\xff')
        obj.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(1),NameObject('/Height'):NumberObject(1),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/XObject'):DictionaryObject({NameObject('/Im1'):writer._add_object(obj)})})
        content=b'q 100 0 0 100 0 0 cm /Im1 Do Q'
    else:
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        content=f'BT /F1 12 Tf 50 780 Td ({text}) Tj ET'.encode()
    stream=DecodedStreamObject();stream.set_data(content);page[NameObject('/Contents')]=writer._add_object(stream)
    if encrypted:writer.encrypt('secret')
    b=BytesIO();writer.write(b);return b.getvalue()


@pytest.mark.parametrize('suffix,encoding',[('.txt','utf-8-sig'),('.md','utf-8'),('.txt','utf-16'),('.txt','gb18030')])
def test_chinese_text_encodings(suffix,encoding):
    text,meta=extract_document('交接'+suffix,SOURCE.encode(encoding))
    assert text==SOURCE
    assert meta['encoding']==('utf-8-sig' if encoding=='utf-8' else encoding)


def test_word_preserves_paragraph_table_order_and_warns_omissions():
    text,meta=extract_document('交接.DOCX',docx_bytes())
    assert text.index('任务交接材料')<text.index('目标\t说明交付物')<text.index(SOURCE)
    assert meta['warnings'] and meta['format']=='docx'


def test_pdf_text_scan_and_password():
    text,meta=extract_document('text.pdf',pdf_bytes())
    assert text=='Document input test' and meta['page_count']==1
    for raw,code in [(pdf_bytes(image=True),'PDF_OCR_REQUIRED'),(pdf_bytes(encrypted=True),'DOCUMENT_ENCRYPTED')]:
        with pytest.raises(DocumentError) as e:extract_document('bad.pdf',raw)
        assert e.value.code==code


def test_invalid_empty_and_unaccepted_word_revisions():
    for filename,raw in [('old.doc',b'doc'),('broken.pdf',b'not pdf'),('empty.txt',b''),('binary.txt',b'hello\x00world'),('broken.docx',b'PKbroken')]:
        with pytest.raises(DocumentError):extract_document(filename,raw)
    original=BytesIO(docx_bytes());edited=BytesIO()
    with zipfile.ZipFile(original) as src, zipfile.ZipFile(edited,'w') as dst:
        for info in src.infolist():
            raw=src.read(info.filename)
            if info.filename=='word/document.xml':raw=raw.replace(b'<w:body>',b'<w:body><w:ins w:id="1"/>')
            dst.writestr(info,raw)
    with pytest.raises(DocumentError) as e:extract_document('revisions.docx',edited.getvalue())
    assert e.value.code=='DOCUMENT_TRACKED_CHANGES'


def test_upload_to_edited_job_survives_restart_and_import_removal(tmp_path):
    settings=Settings(data_dir=tmp_path/'data')
    raw=docx_bytes()
    with TestClient(create_app(settings,start_worker=False,preflight=READY)) as c:
        uploaded=c.post('/api/input-documents',files={'file':('../材料.docx',raw)})
        assert uploaded.status_code==201
        document=uploaded.json();assert document['filename']=='材料.docx'
        updated=document['text']+'\n本次交付物是一页摘要。'
        r=c.post('/api/jobs',json={'source_text':updated,'input_document_id':document['document_id'],'video_form':'auto','duration_mode':'content'})
        assert r.status_code==202
        jid=r.json()['job_id'];store=c.app.state.store
        assert c.get(f'/api/jobs/{jid}/editable-content').json()['source_text']==updated
        info=c.get(f'/api/jobs/{jid}/input-document').json()
        assert info['document']['source_edited']
        assert c.get(info['download_url']).content==raw
        import shutil
        shutil.rmtree(settings.data_dir/'input-documents'/document['document_id'])
    with TestClient(create_app(settings,start_worker=False,preflight=READY)) as c:
        assert c.get(f'/api/jobs/{jid}/input-document/original').content==raw
        path=settings.data_dir/'jobs'/jid/'input/document/original.docx'
        path.write_bytes(b'changed')
        assert c.get(f'/api/jobs/{jid}/input-document/original').status_code==409


def test_bad_upload_does_not_create_job_and_missing_or_changed_import_rejected(tmp_path):
    settings=Settings(data_dir=tmp_path/'data')
    with TestClient(create_app(settings,start_worker=False,preflight=READY)) as c:
        assert c.post('/api/input-documents',files={'file':('bad.pdf',pdf_bytes(image=True))}).status_code==422
        for token in [str(uuid4()),'../../outside']:
            r=c.post('/api/jobs',json={'source_text':SOURCE,'input_document_id':token})
            assert r.status_code==409
        uploaded=c.post('/api/input-documents',files={'file':('材料.txt',SOURCE.encode())}).json()
        (settings.data_dir/'input-documents'/uploaded['document_id']/'extracted.txt').write_text('tampered')
        r=c.post('/api/jobs',json={'source_text':SOURCE,'input_document_id':uploaded['document_id']})
        assert r.status_code==409
        assert c.get('/api/jobs').json()['total']==0


def test_upload_limit_does_not_truncate_and_plain_text_client_compatible(tmp_path,monkeypatch):
    import g3.app as module
    monkeypatch.setattr(module,'MAX_UPLOAD_BYTES',16)
    with TestClient(create_app(Settings(data_dir=tmp_path/'data'),start_worker=False,preflight=READY)) as c:
        assert c.post('/api/input-documents',files={'file':('a.txt',b'a'*17)}).status_code==413
        r=c.post('/api/jobs',json={'source_text':SOURCE})
        assert r.status_code==202
        assert c.get('/api/jobs/'+r.json()['job_id']+'/input-document').json()['document'] is None


def test_document_extraction_available_without_model_configuration(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path/'data'),start_worker=False,preflight=lambda _:{'ready':False})) as c:
        assert c.post('/api/input-documents',files={'file':('材料.md',SOURCE.encode())}).status_code==201
