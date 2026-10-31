"""Article import, fidelity, versioning and shared rendering contracts."""
import json
from io import BytesIO

import pytest

from core.video import article_writer, articles, credibility, store
from core.llm import ChatResponse
from tests.test_video_scripts import video_env, autovideo, _login, _csrf, _wait  # noqa: F401
from web.services import video_executor, video_render, video_workflow
from web.settings import get_settings

PROSE = "雨停以后，窗台还留着一点水光。\n\n我没有急着关窗。风走进来，翻动了桌上那本没有读完的书。"
INFO = "这是一份演示记录：在室温 25 摄氏度、亮度固定的网页浏览测试中，样机连续运行 10 小时。这个成绩不代表日常混合使用，测试也没有比较其他机型。"


class Model:
    calls = []
    result = {"faithful": True, "summary": "保留了测试条件、结果和适用范围", "issues": []}

    def __init__(self, *args, **kwargs):
        self.usage = {"requests": 0}

    def chat(self, messages, **kwargs):
        self.calls.append(messages)
        self.usage["requests"] += 1
        if messages[0]["content"] == article_writer.REVIEW_PROMPT:
            return ChatResponse(content=json.dumps(self.result, ensure_ascii=False))
        return ChatResponse(content=json.dumps({"title": "样机续航十小时，只限这次测试", "sections": [{
            "title": "测试与边界", "paragraphs": [{"text": INFO, "citations": []}]}]}, ensure_ascii=False))


@pytest.fixture(autouse=True)
def reset_model():
    Model.calls = []
    Model.result = {"faithful": True, "summary": "保留了测试条件、结果和适用范围", "issues": []}


def create(client, **values):
    _login(client)
    response = client.post("/api/video/articles/scripts", headers=_csrf(client), json={
        "title": "雨停以后", "text": PROSE, "mode": "read", "kind": "literary", "column": "deep_dive", **values})
    assert response.status_code == 202, response.text
    script_id = response.json()["id"]
    _wait(client, script_id, {"DRAFT", "FAILED"})
    return store.get_script(script_id)


def test_paste_reading_skips_research_and_model_preserves_original(client, video_env, monkeypatch):
    monkeypatch.setattr("core.llm.LLMClient", lambda *args, **kwargs: pytest.fail("原文朗读不应调用模型"))
    script = create(client)
    assert script["status"] == "DRAFT" and not script["report_path"] and not script["research_job"]
    version = store.get_version(script["id"], 1)
    assert article_writer.canonical(article_writer.narration(version["sections"])) == article_writer.canonical(PROSE)
    assert version["meta"]["checks"]["summary"]["passed"]
    assert not version["meta"]["checks"]["external_facts_checked"]
    assert credibility.label(version)["text"] == "已核对原文"
    detail = client.get(f"/api/video/scripts/{script['id']}").json()
    assert detail["steps"][0]["state"] == "SKIPPED"
    assert detail["steps"][2]["state"] == "DONE"
    original = client.get(f"/api/video/scripts/{script['id']}/article").json()
    assert original["text"] == PROSE
    assert client.get(original["original_url"]).content.decode() == PROSE


def test_uploaded_text_and_cleaned_working_copy_are_kept_separately(client, video_env):
    _login(client)
    uploaded = client.post("/api/video/articles/import", headers=_csrf(client),
                           files={"file": ("../文章.txt", ("多余页眉\n\n" + PROSE).encode(), "text/plain")})
    assert uploaded.status_code == 201, uploaded.text
    imported = uploaded.json()
    assert imported["filename"] == "文章.txt"
    script = create(client, import_id=imported["id"])
    original = client.get(f"/api/video/scripts/{script['id']}/article").json()
    assert original["text"] == PROSE and "多余页眉" in original["extracted_text"]
    assert original["edited_after_import"]


@pytest.mark.parametrize("filename,raw", [("bad.exe", b"x"), ("empty.txt", b""),
                                          ("bad.pdf", b"not pdf"), ("bad.docx", b"bad")])
def test_bad_uploads_fail_without_creating_scripts(client, video_env, filename, raw):
    _login(client)
    response = client.post("/api/video/articles/import", headers=_csrf(client), files={"file": (filename, raw)})
    assert response.status_code == 422
    assert not store.list_scripts()


def test_word_extraction_keeps_paragraph_and_table_order():
    from docx import Document
    document = Document(); document.add_paragraph("雨停以后。")
    table = document.add_table(rows=1, cols=2); table.cell(0, 0).text = "条件"; table.cell(0, 1).text = "25 摄氏度"
    document.add_paragraph("最后一段。")
    buffer = BytesIO(); document.save(buffer)
    body, warnings = articles.extract("文章.docx", buffer.getvalue())
    assert body.index("雨停") < body.index("条件") < body.index("最后")
    assert "25 摄氏度" in body and warnings


def test_word_upload_and_original_download_roundtrip(client, video_env):
    from docx import Document
    document = Document()
    document.add_paragraph(PROSE)
    buffer = BytesIO()
    document.save(buffer)
    original_bytes = buffer.getvalue()
    _login(client)
    uploaded = client.post("/api/video/articles/import", headers=_csrf(client),
                           files={"file": ("散文.docx", original_bytes)})
    assert uploaded.status_code == 201
    imported = uploaded.json()
    assert imported["text"] == PROSE
    script = create(client, import_id=imported["id"])
    source = client.get(f"/api/video/scripts/{script['id']}/article").json()
    assert source["filename"] == "散文.docx"
    assert client.get(source["original_url"]).content == original_bytes


def test_text_pdf_is_extracted_without_losing_the_content():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 30 350 Td (Water can evaporate below boiling.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    body, warnings = articles.extract("article.pdf", buffer.getvalue())
    assert body == "Water can evaporate below boiling."
    assert warnings


def test_malformed_source_url_is_a_validation_error_not_server_error(client, video_env):
    _login(client)
    response = client.post("/api/video/articles/scripts", headers=_csrf(client),
                           json={"text": PROSE, "mode": "read", "kind": "literary", "source_url": "https://["})
    assert response.status_code == 422
    assert not store.list_scripts()


def test_snapshot_detects_changed_working_text_and_original(tmp_path):
    source = articles.snapshot(str(tmp_path), body=PROSE, title="雨停以后")
    original = articles.read(str(tmp_path), source)
    path = tmp_path / "video_articles" / source["id"] / "working.txt"
    path.write_text("偷偷改写")
    with pytest.raises(articles.ArticleError, match="哈希"):
        articles.read(str(tmp_path), source)
    path.write_text(PROSE)
    raw = articles.read_import(str(tmp_path), original["import_id"])
    from pathlib import Path
    Path(raw["original_path"]).write_text("原件被修改")
    with pytest.raises(articles.ArticleError):
        articles.read(str(tmp_path), source)


def test_source_metadata_cannot_change_independently_of_saved_reference(tmp_path):
    source = articles.snapshot(str(tmp_path), body=PROSE, title="雨停以后")
    path = tmp_path / "video_articles" / source["id"] / "source.json"
    altered = {**source, "source_url": "https://example.com/new-source"}
    path.write_text(json.dumps(altered))
    with pytest.raises(articles.ArticleError, match="哈希"):
        articles.read(str(tmp_path), source)


def test_scan_pdf_is_explicitly_rejected_without_importing_partial_text():
    from PIL import Image
    buffer = BytesIO(); Image.new("RGB", (100, 100), "white").save(buffer, format="PDF")
    with pytest.raises(articles.ArticleError, match="OCR"):
        articles.extract("扫描.pdf", buffer.getvalue())


def test_adaptation_audits_against_original_without_claiming_external_verification(client, video_env, monkeypatch):
    monkeypatch.setattr("core.llm.LLMClient", Model)
    script = create(client, title="续航测试", text=INFO, mode="adapt", kind="information", origin="ai")
    assert script["status"] == "DRAFT"
    version = store.get_version(script["id"], 1)
    assert version["meta"]["checks"]["summary"]["passed"]
    assert len(Model.calls) == 2
    assert INFO in Model.calls[-1][1]["content"]
    assert "未做外部事实核验" in "".join(credibility.label(version)["reasons"])


def test_unfounded_numbers_and_lost_conditions_block_approval(client, video_env, monkeypatch):
    monkeypatch.setattr("core.llm.LLMClient", Model)
    Model.result = {"faithful": False, "summary": "删掉了测试条件", "issues": [{
        "section": 0, "paragraph": 0, "quote": "样机连续运行 10 小时", "severity": "must",
        "problem": "结论适用范围丢失", "suggestion": "保留测试条件"}]}
    script = create(client, text=INFO, mode="adapt", kind="information")
    with pytest.raises(video_executor.StepError, match="忠实度问题"):
        video_executor.approve(script["id"], 1, "test")
    assert article_writer.local_issues([{"paragraphs": [{"text": "续航 99 小时。"}]}], {"text": INFO}, "adapt", "information")


def test_invalid_audit_never_looks_like_pass(client, video_env, monkeypatch):
    monkeypatch.setattr("core.llm.LLMClient", Model)
    Model.result = {"faithful": True, "summary": "通过", "issues": [{"section": 0, "paragraph": 0,
        "quote": "不存在的引文", "severity": "warn", "problem": "建议", "suggestion": ""}]}
    script = create(client, text=INFO, mode="adapt", kind="information")
    assert script["error"] and not script["stage"]
    with pytest.raises(video_executor.StepError, match="尚未完成"):
        video_executor.approve(script["id"], 1, "test")


def test_edit_invalidates_checks_and_rerun_checks_current_version(client, video_env):
    script = create(client)
    old = store.get_version(script["id"], 1)
    n = video_workflow.save_edit(script["id"], title=old["title"], sections=old["sections"], edited_by="test")
    with pytest.raises(video_executor.StepError, match="尚未完成"):
        video_executor.approve(script["id"], n, "test")
    from web.services.video_articles import audit
    audit(store.get_script(script["id"]))
    video_executor.approve(script["id"], n, "test")
    assert store.get_version(script["id"], 1) == old


def test_article_uses_shared_render_with_literary_intent_and_no_fake_sources(client, video_env, autovideo):
    script = create(client)
    video_executor.approve(script["id"], 1, "test")
    render = video_render.create(store.get_script(script["id"]))
    request = render["request"]
    assert "sources" not in request and request["input_mode"] == "narration"
    assert "文学朗读" in request["creative_intent"]["purpose"]
    assert request["packaging"]["label"] == "文章朗读"
    assert video_render.process(render["id"]) == "COMPLETED"
    assert autovideo.posts[-1]["chapters"] == request["chapters"]


def test_matrix_partial_packaging_inherits_recipe_and_removes_unbacked_source_hint(client, video_env, monkeypatch):
    from dataclasses import replace
    from types import SimpleNamespace
    from core.video import columns
    column_get = columns.get
    column = replace(column_get(str(video_env / "config"), "deep_dive"),
                     packaging={"preset": "common", "label": "原有栏目", "outro": {"lines": ["主要来源见简介"], "seconds": 3}})
    monkeypatch.setattr(columns, "get", lambda *args, **kwargs: column)
    config = SimpleNamespace(recipe="deep_dive", key="linked", usable=True, job="",
                             render={"video_speed": 1.2, "packaging": {"sound": {"enabled": False}}})
    monkeypatch.setattr("core.video.production.get", lambda *args, **kwargs: config)
    script = create(client, production="linked")
    assert script["production"] == "linked"
    video_executor.approve(script["id"], 1, "test")
    request = video_render.create(store.get_script(script["id"]))["request"]
    assert request["production_context"]["matrix"]["settings"]["video_speed"] == 1.2
    assert request["packaging"]["outro"]["lines"] == ["感谢收听"]
    assert not request["packaging"]["sound"]["enabled"]
    assert request["packaging"]["label"] == "原有栏目"


def test_dated_article_source_uses_existing_source_card_contract(client, video_env):
    script = create(client, source_url="https://example.com/article", published_on="2026-10-25")
    video_executor.approve(script["id"], 1, "test")
    request = video_render.create(store.get_script(script["id"]))["request"]
    assert request["sources"] == [{"title": "雨停以后", "url": "https://example.com/article", "date": "2026-10-25"}]


def test_imported_article_can_find_visuals_from_its_link_without_a_research_report(client, video_env, monkeypatch):
    url = "https://playwright.dev/python/docs/screenshots"
    script = create(client, title="Playwright 元素截图", text="Playwright 用 locator 选择元素，再保存截图。",
                    mode="read", kind="information", source_url=url)
    store.update_script(script["id"], status="APPROVED", approved_version=1)
    monkeypatch.setattr(video_render, "_video_config", lambda: {"visual_materials": {"enabled": True, "version": 3}})
    assert not script["report_path"]
    render = video_render.create(store.get_script(script["id"]))
    assert render["visuals"]["state"] == "pending"
    assert render["visuals"]["needs"][0]["urls"] == [url]


@pytest.mark.parametrize("extra", [{"source_url": "javascript:alert(1)"}, {"published_on": "2026-02-30"},
                                  {"import_id": "../../.env"}, {"mode": "bad"}])
def test_invalid_article_input_is_rejected(client, video_env, extra):
    _login(client)
    response = client.post("/api/video/articles/scripts", headers=_csrf(client), json={
        "text": PROSE, "mode": "read", "kind": "literary", "column": "deep_dive", **extra})
    assert response.status_code == 422
    assert not store.list_scripts()
